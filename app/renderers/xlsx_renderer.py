from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)

from ..output_models import GeneratedDocument


INVALID_SHEET_CHARS = set(r'[]:*?/\\')


def _safe_sheet_name(name: str) -> str:
    cleaned = "".join(
        "-" if char in INVALID_SHEET_CHARS else char
        for char in str(name)
    )

    cleaned = cleaned.strip() or "Study Notes"

    return cleaned[:31]


def _unique_sheet_name(wb, name: str) -> str:
    original = _safe_sheet_name(name)
    candidate = original
    suffix = 2

    while candidate in wb.sheetnames:
        tail = f"-{suffix}"
        candidate = (
            original[:31 - len(tail)] + tail
        )
        suffix += 1

    return candidate


def _safe_table_name(name: str, index: int) -> str:
    cleaned = "".join(
        char if char.isalnum() else "_"
        for char in str(name)
    ).strip("_")

    if not cleaned:
        cleaned = f"StudyTable_{index}"

    if cleaned[0].isdigit():
        cleaned = "T_" + cleaned

    return (
        f"{cleaned[:220]}_{index}"
    )


def _estimate_column_width(values) -> float:
    lengths = []

    for value in values:
        if value is None:
            continue

        text = str(value)

        # Long academic text should wrap instead of producing
        # absurdly wide sheets.
        longest_line = max(
            (len(line) for line in text.splitlines()),
            default=0,
        )

        lengths.append(longest_line)

    if not lengths:
        return 16

    return min(
        max(max(lengths) + 2, 12),
        42,
    )


def _style_overview(ws):
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 90

    ws.freeze_panes = "A4"


def _style_table_sheet(
    ws,
    headers,
    rows,
    table_index,
):
    ws.sheet_view.showGridLines = False

    if headers:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}"
            f"{max(len(rows) + 1, 1)}"
        )

    max_columns = max(
        [len(headers)]
        + [len(row) for row in rows]
        + [1]
    )

    for column in range(
        1,
        max_columns + 1,
    ):
        values = []

        if column <= len(headers):
            values.append(
                headers[column - 1]
            )

        for row in rows:
            if column <= len(row):
                values.append(
                    row[column - 1]
                )

        ws.column_dimensions[
            get_column_letter(column)
        ].width = _estimate_column_width(
            values
        )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    if headers:
        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="315B49",
            )

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="center",
            )

        ws.row_dimensions[1].height = 32

    for row_number in range(
        2,
        len(rows) + 2,
    ):
        ws.row_dimensions[
            row_number
        ].height = 42

    # Native Excel table gives filtering, banding and better
    # navigation without asking the AI to design formatting.
    if headers and rows:
        end_column = get_column_letter(
            len(headers)
        )

        end_row = len(rows) + 1

        ref = (
            f"A1:{end_column}{end_row}"
        )

        table = Table(
            displayName=_safe_table_name(
                ws.title,
                table_index,
            ),
            ref=ref,
        )

        style = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = style

        ws.add_table(table)


def _add_navigation_sheet(
    wb,
    spec,
    table_sheets,
):
    ws = wb["Overview"]

    row = 4

    ws.cell(
        row=row,
        column=1,
        value="Workbook contents",
    ).font = Font(
        bold=True,
        size=13,
    )

    row += 1

    for sheet_name, description in table_sheets:
        cell = ws.cell(
            row=row,
            column=1,
            value=sheet_name,
        )

        cell.hyperlink = (
            f"#'{sheet_name}'!A1"
        )

        cell.style = "Hyperlink"

        ws.cell(
            row=row,
            column=2,
            value=description,
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        row += 1

    if table_sheets:
        row += 1

    return row


def render_xlsx(
    spec: GeneratedDocument,
    output_path: Path,
) -> Path:

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb = Workbook()

    overview = wb.active
    overview.title = "Overview"

    _style_overview(overview)

    overview["A1"] = spec.title
    overview["A1"].font = Font(
        size=20,
        bold=True,
    )

    if spec.subtitle:
        overview["A2"] = spec.subtitle
        overview["A2"].font = Font(
            italic=True,
        )

    # --------------------------------------------------------
    # Create structured worksheets first
    # --------------------------------------------------------

    table_sheets = []
    table_counter = 1

    for section in spec.sections:
        for table_index, table_spec in enumerate(
            section.tables,
            start=1,
        ):
            title = (
                table_spec.title
                or (
                    section.heading
                    if len(section.tables) == 1
                    else (
                        f"{section.heading} "
                        f"{table_index}"
                    )
                )
            )

            sheet_name = _unique_sheet_name(
                wb,
                title,
            )

            ws = wb.create_sheet(
                sheet_name
            )

            headers = [
                str(value)
                for value in table_spec.headers
            ]

            rows = [
                [
                    (
                        value
                        if value is not None
                        else ""
                    )
                    for value in row
                ]
                for row in table_spec.rows
            ]

            # Ensure rows never exceed header width when headers
            # exist; extra values would otherwise make malformed
            # native Excel tables.
            if headers:
                normalised_rows = []

                for row in rows:
                    if len(row) < len(headers):
                        row = row + (
                            [""] * (
                                len(headers)
                                - len(row)
                            )
                        )

                    elif len(row) > len(headers):
                        row = row[
                            :len(headers)
                        ]

                    normalised_rows.append(
                        row
                    )

                rows = normalised_rows

            if headers:
                for column, value in enumerate(
                    headers,
                    start=1,
                ):
                    ws.cell(
                        row=1,
                        column=column,
                        value=value,
                    )

            start_row = (
                2
                if headers
                else 1
            )

            for row_number, values in enumerate(
                rows,
                start=start_row,
            ):
                for column, value in enumerate(
                    values,
                    start=1,
                ):
                    ws.cell(
                        row=row_number,
                        column=column,
                        value=value,
                    )

            _style_table_sheet(
                ws,
                headers,
                rows,
                table_counter,
            )

            description = (
                section.takeaway
                or section.slide_purpose
                or section.heading
            )

            table_sheets.append(
                (
                    sheet_name,
                    description,
                )
            )

            table_counter += 1

    # --------------------------------------------------------
    # Overview + navigation + narrative context
    # --------------------------------------------------------

    row = _add_navigation_sheet(
        wb,
        spec,
        table_sheets,
    )

    for section in spec.sections:
        if (
            not section.paragraphs
            and not section.bullets
        ):
            continue

        overview.cell(
            row=row,
            column=1,
            value=section.heading,
        )

        overview.cell(
            row=row,
            column=1,
        ).font = Font(
            bold=True,
            size=13,
        )

        row += 1

        for paragraph in section.paragraphs:
            overview.cell(
                row=row,
                column=2,
                value=paragraph,
            ).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

            row += 1

        for bullet in section.bullets:
            overview.cell(
                row=row,
                column=1,
                value="•",
            )

            overview.cell(
                row=row,
                column=2,
                value=bullet,
            ).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

            row += 1

        row += 1

    # --------------------------------------------------------
    # References
    # --------------------------------------------------------

    if spec.references:
        refs = wb.create_sheet(
            _unique_sheet_name(
                wb,
                "References",
            )
        )

        refs.sheet_view.showGridLines = False

        refs["A1"] = "APA 7 References"

        refs["A1"].font = Font(
            bold=True,
            size=15,
        )

        refs.column_dimensions[
            "A"
        ].width = 110

        for index, reference in enumerate(
            spec.references,
            start=3,
        ):
            cell = refs.cell(
                row=index,
                column=1,
                value=reference,
            )

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

            refs.row_dimensions[
                index
            ].height = 35

    # --------------------------------------------------------
    # Consistent view
    # --------------------------------------------------------

    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90

    wb.save(output_path)

    return output_path
