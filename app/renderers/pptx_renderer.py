from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ..output_models import GeneratedDocument


def _add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    bold=False,
):
    box = slide.shapes.add_textbox(
        left,
        top,
        width,
        height,
    )

    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True

    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold

    return box


def _add_bullets(
    slide,
    items,
    left,
    top,
    width,
    height,
    size=19,
):
    box = slide.shapes.add_textbox(
        left,
        top,
        width,
        height,
    )

    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True

    for index, item in enumerate(items):
        paragraph = (
            frame.paragraphs[0]
            if index == 0
            else frame.add_paragraph()
        )

        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(size)
        paragraph.space_after = Pt(10)

    return box


def _add_table(slide, table_spec):
    headers = table_spec.headers or []

    rows = table_spec.rows or []

    if not headers or not rows:
        return

    row_count = len(rows) + 1
    column_count = len(headers)

    left = Inches(0.7)
    top = Inches(1.75)
    width = Inches(11.9)
    height = Inches(
        min(4.6, 0.55 + (row_count * 0.52))
    )

    shape = slide.shapes.add_table(
        row_count,
        column_count,
        left,
        top,
        width,
        height,
    )

    table = shape.table

    for column, header in enumerate(headers):
        cell = table.cell(0, column)
        cell.text = str(header)

        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(14)

    for row_index, row in enumerate(rows, start=1):
        for column in range(column_count):
            value = (
                row[column]
                if column < len(row)
                else ""
            )

            cell = table.cell(
                row_index,
                column,
            )

            cell.text = str(value)

            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(12)


def _add_takeaway(slide, takeaway):
    if not takeaway:
        return

    _add_textbox(
        slide,
        Inches(0.8),
        Inches(6.35),
        Inches(11.6),
        Inches(0.55),
        f"Key takeaway: {takeaway}",
        size=13,
        bold=True,
    )


def _add_visual_suggestion(slide, suggestion):
    """
    We deliberately do not invent/download images.

    Instead, preserve the AI's meaningful visual design recommendation
    in a subtle presenter-facing placeholder. Later this can feed the
    StudyMate visual-generation pipeline.
    """

    if not suggestion:
        return

    box = _add_textbox(
        slide,
        Inches(8.65),
        Inches(5.55),
        Inches(3.7),
        Inches(0.55),
        f"Visual idea: {suggestion}",
        size=9,
    )

    frame = box.text_frame

    for paragraph in frame.paragraphs:
        paragraph.alignment = PP_ALIGN.RIGHT


def _add_speaker_notes(slide, notes):
    """
    python-pptx exposes notes text frames on supported versions.
    If unavailable, generation should still succeed.
    """

    if not notes:
        return

    try:
        notes_slide = slide.notes_slide

        text_frame = notes_slide.notes_text_frame

        if text_frame is None:
            return

        existing = text_frame.text.strip()

        content = "\n".join(notes)

        if existing:
            text_frame.text = (
                existing
                + "\n\nEl Amigo speaker notes:\n"
                + content
            )
        else:
            text_frame.text = (
                "El Amigo speaker notes:\n"
                + content
            )

    except Exception:
        # Speaker notes must never prevent the actual deck rendering.
        pass


def _render_content_slide(prs, section):
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )

    slide.shapes.title.text = section.heading

    content = []

    content.extend(section.paragraphs)
    content.extend(section.bullets)

    # Presentation specialist should keep this concise,
    # but enforce a final safety limit in the renderer.
    content = content[:6]

    if content:
        _add_bullets(
            slide,
            content,
            Inches(0.85),
            Inches(1.65),
            Inches(11.4),
            Inches(4.45),
            size=19,
        )

    _add_takeaway(
        slide,
        section.takeaway,
    )

    _add_visual_suggestion(
        slide,
        section.visual_suggestion,
    )

    _add_speaker_notes(
        slide,
        section.speaker_notes,
    )


def _render_comparison_slide(prs, section):
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )

    slide.shapes.title.text = section.heading

    if section.tables:
        _add_table(
            slide,
            section.tables[0],
        )

    else:
        items = (
            section.bullets
            or section.paragraphs
        )[:8]

        midpoint = max(
            1,
            (len(items) + 1) // 2,
        )

        left_items = items[:midpoint]
        right_items = items[midpoint:]

        _add_bullets(
            slide,
            left_items,
            Inches(0.75),
            Inches(1.7),
            Inches(5.65),
            Inches(4.3),
            size=18,
        )

        if right_items:
            _add_bullets(
                slide,
                right_items,
                Inches(6.85),
                Inches(1.7),
                Inches(5.65),
                Inches(4.3),
                size=18,
            )

    _add_takeaway(
        slide,
        section.takeaway,
    )

    _add_speaker_notes(
        slide,
        section.speaker_notes,
    )


def _render_table_slide(prs, section):
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )

    slide.shapes.title.text = section.heading

    if section.tables:
        table = section.tables[0]

        if table.title:
            _add_textbox(
                slide,
                Inches(0.75),
                Inches(1.25),
                Inches(11.5),
                Inches(0.35),
                table.title,
                size=12,
                bold=True,
            )

        _add_table(
            slide,
            table,
        )

    else:
        _add_bullets(
            slide,
            (
                section.bullets
                or section.paragraphs
            )[:6],
            Inches(0.85),
            Inches(1.7),
            Inches(11.4),
            Inches(4.3),
            size=18,
        )

    _add_takeaway(
        slide,
        section.takeaway,
    )

    _add_speaker_notes(
        slide,
        section.speaker_notes,
    )


def _render_key_point_slide(prs, section):
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )

    slide.shapes.title.text = section.heading

    key_text = (
        section.takeaway
        or (
            section.paragraphs[0]
            if section.paragraphs
            else (
                section.bullets[0]
                if section.bullets
                else ""
            )
        )
    )

    box = _add_textbox(
        slide,
        Inches(1.5),
        Inches(2.25),
        Inches(10.3),
        Inches(2.3),
        key_text,
        size=28,
        bold=True,
    )

    for paragraph in box.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER

    _add_speaker_notes(
        slide,
        section.speaker_notes,
    )


def _render_process_slide(prs, section):
    slide = prs.slides.add_slide(
        prs.slide_layouts[5]
    )

    slide.shapes.title.text = section.heading

    steps = (
        section.bullets
        or section.paragraphs
    )[:6]

    if not steps:
        steps = ["Process"]

    usable_width = 11.5
    gap = 0.15

    step_width = (
        usable_width
        - gap * (len(steps) - 1)
    ) / len(steps)

    for index, step in enumerate(steps):
        left = (
            0.75
            + index * (step_width + gap)
        )

        _add_textbox(
            slide,
            Inches(left),
            Inches(2.25),
            Inches(step_width),
            Inches(2.2),
            f"{index + 1}\n{step}",
            size=15,
            bold=index == 0,
        )

    _add_takeaway(
        slide,
        section.takeaway,
    )

    _add_speaker_notes(
        slide,
        section.speaker_notes,
    )


def render_pptx(
    spec: GeneratedDocument,
    output_path: Path,
) -> Path:

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    prs = Presentation()

    # 16:9 widescreen.
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --------------------------------------------------------
    # Title slide
    # --------------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[0]
    )

    slide.shapes.title.text = spec.title

    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = (
            spec.subtitle
            or "El Amigo · StudyMate"
        )

    # --------------------------------------------------------
    # AI-designed slides
    # --------------------------------------------------------

    for section in spec.sections:

        layout = (
            section.slide_layout
            or (
                "table"
                if section.tables
                else "content"
            )
        ).lower()

        if layout == "comparison":
            _render_comparison_slide(
                prs,
                section,
            )

        elif layout == "table":
            _render_table_slide(
                prs,
                section,
            )

        elif layout == "process":
            _render_process_slide(
                prs,
                section,
            )

        elif layout == "key_point":
            _render_key_point_slide(
                prs,
                section,
            )

        else:
            _render_content_slide(
                prs,
                section,
            )

    # --------------------------------------------------------
    # References
    # --------------------------------------------------------

    if spec.references:
        chunks = [
            spec.references[i:i + 5]
            for i in range(
                0,
                len(spec.references),
                5,
            )
        ]

        for index, refs in enumerate(chunks):
            slide = prs.slides.add_slide(
                prs.slide_layouts[5]
            )

            title = "References"

            if len(chunks) > 1:
                title += (
                    f" ({index + 1}/{len(chunks)})"
                )

            slide.shapes.title.text = title

            _add_bullets(
                slide,
                refs,
                Inches(0.8),
                Inches(1.6),
                Inches(11.7),
                Inches(5.2),
                size=13,
            )

    prs.save(output_path)

    return output_path
