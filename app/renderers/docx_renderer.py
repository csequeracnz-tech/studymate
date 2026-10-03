from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from ..output_models import GeneratedDocument


def _shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)

    tc_pr.append(shading)


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")

    tr_pr.append(tbl_header)


def _configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)

    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(26)
    title.font.bold = True

    for name, size in (
        ("Heading 1", 17),
        ("Heading 2", 14),
        ("Heading 3", 12),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True


def _add_table(doc, table_spec):
    if table_spec.title:
        p = doc.add_paragraph()

        run = p.add_run(
            str(table_spec.title)
        )

        run.bold = True

        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(5)

    rows = len(table_spec.rows) + (
        1 if table_spec.headers else 0
    )

    cols = max(
        [len(table_spec.headers)]
        + [len(row) for row in table_spec.rows]
        + [1]
    )

    table = doc.add_table(
        rows=max(rows, 1),
        cols=cols,
    )

    table.style = "Table Grid"
    table.autofit = True

    row_index = 0

    if table_spec.headers:
        header_row = table.rows[0]

        for col_index in range(cols):
            value = (
                table_spec.headers[col_index]
                if col_index < len(table_spec.headers)
                else ""
            )

            cell = header_row.cells[col_index]
            cell.text = str(value)
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            _shade_cell(
                cell,
                "E7F0EB",
            )

            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)

        _set_repeat_table_header(
            header_row
        )

        row_index = 1

    for values in table_spec.rows:
        for col_index in range(cols):
            value = (
                values[col_index]
                if col_index < len(values)
                else ""
            )

            cell = table.cell(
                row_index,
                col_index,
            )

            cell.text = str(value)

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.TOP
            )

            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)

        row_index += 1

    doc.add_paragraph()


def _add_takeaway(doc, text):
    if not text:
        return

    p = doc.add_paragraph()

    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)

    run = p.add_run(
        f"Key takeaway: {text}"
    )

    run.bold = True


def render_docx(
    spec: GeneratedDocument,
    output_path: Path,
) -> Path:

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    _configure_styles(doc)

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    title = doc.add_paragraph(
        style="Title"
    )

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title.add_run(spec.title)

    if spec.subtitle:
        subtitle = doc.add_paragraph()

        subtitle.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = subtitle.add_run(
            spec.subtitle
        )

        run.italic = True
        run.font.size = Pt(11)

    doc.add_paragraph()

    # --------------------------------------------------------
    # Academic content
    # --------------------------------------------------------

    for item in spec.sections:
        doc.add_heading(
            item.heading,
            level=1,
        )

        for text in item.paragraphs:
            paragraph = doc.add_paragraph(
                text
            )

            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(8)

        for bullet in item.bullets:
            paragraph = doc.add_paragraph(
                bullet,
                style="List Bullet",
            )

            paragraph.paragraph_format.space_after = Pt(4)

        for table_spec in item.tables:
            _add_table(
                doc,
                table_spec,
            )

        _add_takeaway(
            doc,
            item.takeaway,
        )

    # --------------------------------------------------------
    # References
    # --------------------------------------------------------

    if spec.references:
        doc.add_page_break()

        doc.add_heading(
            "References",
            level=1,
        )

        for reference in spec.references:
            paragraph = doc.add_paragraph(
                reference
            )

            paragraph.paragraph_format.left_indent = (
                Inches(0.5)
            )

            paragraph.paragraph_format.first_line_indent = (
                Inches(-0.5)
            )

            paragraph.paragraph_format.space_after = (
                Pt(8)
            )

    # --------------------------------------------------------
    # Document metadata
    # --------------------------------------------------------

    properties = doc.core_properties
    properties.title = spec.title
    properties.subject = (
        spec.subtitle
        or "StudyMate academic output"
    )
    properties.author = "El Amigo · StudyMate"

    doc.save(output_path)

    return output_path
