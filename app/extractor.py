from pathlib import Path
import json

from PIL import Image
import pytesseract
from pypdf import PdfReader
from docx import Document as WordDocument
from pptx import Presentation
from openpyxl import load_workbook


def extract_text_file(path: Path) -> dict:
    return {
        "type": "text",
        "content": path.read_text(errors="replace"),
    }


def extract_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""

        pages.append({
            "page": page_number,
            "text": text.strip(),
        })

    return {
        "type": "pdf",
        "pages": pages,
    }


def extract_docx(path: Path) -> dict:
    document = WordDocument(str(path))

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    tables = []

    for table in document.tables:
        rows = []

        for row in table.rows:
            rows.append([
                cell.text.strip()
                for cell in row.cells
            ])

        tables.append(rows)

    return {
        "type": "docx",
        "paragraphs": paragraphs,
        "tables": tables,
    }


def extract_pptx(path: Path) -> dict:
    presentation = Presentation(str(path))
    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text.strip())

        slides.append({
            "slide": slide_number,
            "text": text_parts,
        })

    return {
        "type": "pptx",
        "slides": slides,
    }


def extract_xlsx(path: Path) -> dict:
    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )

    sheets = {}

    for worksheet in workbook.worksheets:
        rows = []

        for row in worksheet.iter_rows(values_only=True):
            rows.append([
                value if value is not None else ""
                for value in row
            ])

        sheets[worksheet.title] = rows

    return {
        "type": "xlsx",
        "sheets": sheets,
    }


def extract_image(path: Path) -> dict:
    image = Image.open(path)

    text = pytesseract.image_to_string(image)

    return {
        "type": "image",
        "width": image.width,
        "height": image.height,
        "ocr_text": text.strip(),
    }


def extract_document(path: str | Path) -> dict:
    path = Path(path)
    extension = path.suffix.lower()

    if extension in {".txt", ".md"}:
        return extract_text_file(path)

    if extension == ".pdf":
        return extract_pdf(path)

    if extension == ".docx":
        return extract_docx(path)

    if extension == ".pptx":
        return extract_pptx(path)

    if extension == ".xlsx":
        return extract_xlsx(path)

    if extension in {".jpg", ".jpeg", ".png", ".webp", ".tiff"}:
        return extract_image(path)

    raise ValueError(
        f"Unsupported extraction type: {extension}"
    )


def save_extraction(path: str | Path) -> Path:
    path = Path(path)

    extracted = extract_document(path)

    output_path = path.with_suffix(".extracted.json")

    output_path.write_text(
        json.dumps(
            extracted,
            indent=2,
            ensure_ascii=False,
            default=str,
        )
    )

    return output_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python -m app.extractor <file>")
        raise SystemExit(1)

    source = Path(sys.argv[1])

    output = save_extraction(source)

    print(f"Extracted: {source}")
    print(f"Saved:     {output}")
